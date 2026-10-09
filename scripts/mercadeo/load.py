"""Load Mercadeo facts from the 9 October 2026 files and the Pipedrive extract.

Run from the repo root, after migration 079 is applied and schema mercadeo is
exposed in the Data API:

    python3 scripts/mercadeo/load.py

    python3 scripts/mercadeo/load.py --check

--check parses the files and prints counts. It does not write.

The load replaces ad_day, budget_month, and deal_attribution. It does not
change cost_center or lead_range. It does not write schema ventas.
"""

from __future__ import annotations

import csv
import json
import os
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import uuid
from decimal import Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "docs-mercadeo-oct-09"
PIPEDRIVE = Path(
    "/Users/orion-tech/go/singularity/_ Agents_Extract_PipeDrive/"
    "repo-orion-odoo19-etl-pipedrive/output/resolved"
)
EXCLUDED = "cons_guat__kpa_LeadsWebsiteBoulevard_fb_lds_cpa_11noval20nov"
BATCH = 80
MONTHS = {
    "ENERO": 1,
    "FEBRERO": 2,
    "MARZO": 3,
    "ABRIL": 4,
    "MAYO": 5,
    "JUNIO": 6,
    "JULIO": 7,
    "AGOSTO": 8,
    "SEPTIEMBRE": 9,
    "OCTUBRE": 10,
    "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}
META_ACCOUNTS = {
    "Benestare": "benestare",
    "Bosque Las Tapias": "bosque-las-tapias",
    "Boulevard5": "boulevard-5",
    "Santa Elena": "santa-elena",
    "Casa Elisa": "casa-elisa",
    "Puerta Abierta": "puerta-abierta",
}
GOOGLE_ACCOUNTS = {
    "Benestare Zona 6": "benestare",
    "Bosque Las Tapias | Apartamentos Zona 18": "bosque-las-tapias",
    "Boulevard 5 GA": "boulevard-5",
    "Santa Elena": "santa-elena",
    "Casa Elisa": "casa-elisa",
}
BUDGET_SHEETS = {
    "Benestare": "benestare",
    "Bosque Las Tapias": "bosque-las-tapias",
    "Boulevard 5": "boulevard-5",
    "Santa Elena": "santa-elena",
    "Puerta Abierta": "puerta-abierta",
}


def uuid7() -> str:
    unix_ms = int(time.time() * 1000) & ((1 << 48) - 1)
    raw = unix_ms.to_bytes(6, "big") + os.urandom(10)
    data = bytearray(raw)
    data[6] = (data[6] & 0x0F) | 0x70
    data[8] = (data[8] & 0x3F) | 0x80
    return str(uuid.UUID(bytes=bytes(data)))


def load_env() -> dict[str, str]:
    found: dict[str, str] = {}
    with open(ROOT / ".env.local", encoding="utf-8") as handle:
        for line in handle:
            text = line.strip()
            if not text or text.startswith("#") or "=" not in text:
                continue
            key, value = text.split("=", 1)
            found[key.strip()] = value.strip().strip('"').strip("'")
    return found


def number_text(value: object) -> str:
    return str(value).strip().replace(",", "").replace(" ", "")


def money(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    amount = Decimal(number_text(value))
    if amount < 0:
        raise RuntimeError(f"negative amount {value}")
    return amount.quantize(Decimal("0.01"))


def qty(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0.00")
    return Decimal(number_text(value)).quantize(Decimal("0.01"))


def plain(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def fold(value: str) -> str:
    text = unicodedata.normalize("NFD", value.strip().lower())
    return "".join(char for char in text if unicodedata.category(char) != "Mn")


def line_code(label: str) -> str | None:
    text = fold(label)
    if text.startswith("meta"):
        return "meta"
    if text.startswith("google"):
        return "google"
    if text.startswith("tik"):
        return "tiktok"
    if text.startswith("linkedin"):
        return "linkedin"
    if text.startswith("pauta digital"):
        return "pauta_digital"
    if text.startswith("wati"):
        return "wati"
    return None


def find_file(prefix: str) -> Path:
    matches = sorted(path for path in DATA.iterdir() if path.name.startswith(prefix))
    if len(matches) != 1:
        raise RuntimeError(f"expected one {prefix} file in {DATA}, found {matches}")
    return matches[0]


class Api:
    def __init__(self, base_url: str, service_key: str) -> None:
        self.base = base_url.rstrip("/") + "/rest/v1"
        self.key = service_key

    def request(
        self,
        method: str,
        path: str,
        payload: object | None = None,
        range_header: str | None = None,
        prefer: str | None = None,
    ) -> tuple[int, dict[str, str], object]:
        headers = {
            "apikey": self.key,
            "Authorization": f"Bearer {self.key}",
            "Accept": "application/json",
            "Content-Profile": "mercadeo",
            "Accept-Profile": "mercadeo",
            "Prefer": prefer or "return=minimal",
        }
        if range_header:
            headers["Range"] = range_header
        data = None
        if payload is not None:
            data = json.dumps(payload).encode("utf-8")
            headers["Content-Type"] = "application/json"
        request = urllib.request.Request(self.base + path, data=data, headers=headers, method=method)
        for attempt in range(6):
            try:
                with urllib.request.urlopen(request, timeout=180) as response:
                    raw = response.read()
                    body: object = json.loads(raw.decode("utf-8")) if raw else None
                    return response.status, dict(response.headers), body
            except urllib.error.HTTPError as error:
                detail = error.read().decode("utf-8", errors="replace")
                if error.code == 416:
                    return 416, dict(error.headers), []
                if error.code in (429, 500, 502, 503, 504) and attempt < 5:
                    time.sleep(2**attempt)
                    continue
                raise RuntimeError(f"{method} {path} failed ({error.code}): {detail[:800]}") from error
        raise RuntimeError(f"{method} {path} failed after retries")

    def fetch_all(self, table: str, select: str) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        start = 0
        size = 1000
        quoted = urllib.parse.quote(select, safe=",")
        while True:
            status, _, body = self.request(
                "GET",
                f"/{table}?select={quoted}&order=id.asc",
                range_header=f"{start}-{start + size - 1}",
                prefer="count=exact",
            )
            if status == 416 or not body:
                break
            if not isinstance(body, list):
                raise RuntimeError(f"expected a list from {table}")
            rows.extend(body)
            if len(body) < size:
                break
            start += size
        return rows

    def insert(self, table: str, rows: list[dict[str, object]]) -> None:
        for offset in range(0, len(rows), BATCH):
            chunk = rows[offset : offset + BATCH]
            self.request("POST", f"/{table}", chunk, prefer="return=minimal")
            done = min(offset + BATCH, len(rows))
            if done == len(rows) or done % (BATCH * 25) == 0:
                print(f"  {table}: {done}/{len(rows)}", flush=True)

    def delete_all(self, table: str) -> None:
        self.request("DELETE", f"/{table}?id=not.is.null", prefer="return=minimal")


def add_ad(bucket: dict[tuple[str, str, str, str], dict[str, object]], row: dict[str, object]) -> None:
    key = (str(row["platform"]), str(row["cost_center_slug"]), str(row["campaign_name"]), str(row["day"]))
    current = bucket.get(key)
    if current is None:
        bucket[key] = row
        return
    if current["currency"] != row["currency"]:
        raise RuntimeError(f"two currencies for {key}")
    for field in ("spend", "impressions", "clicks", "leads", "lead_forms"):
        current[field] = Decimal(str(current[field])) + Decimal(str(row[field]))
    if row["reach"] is not None:
        current["reach"] = Decimal(str(current["reach"] or 0)) + Decimal(str(row["reach"]))


def load_meta() -> dict[tuple[str, str, str, str], dict[str, object]]:
    path = find_file("Performance-Report")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [plain(cell) for cell in next(rows)]
    index = {name: position for position, name in enumerate(header)}
    required = ["Día", "Nombre de la cuenta", "Nombre de la campaña", "Alcance", "Impresiones", "Clientes potenciales", "Divisa", "Clics (todos)", "Importe gastado (USD)"]
    missing = [name for name in required if name not in index]
    if missing:
        raise RuntimeError(f"Meta file is missing {missing}")
    bucket: dict[tuple[str, str, str, str], dict[str, object]] = {}
    unknown: set[str] = set()
    for raw in rows:
        if not raw or raw[index["Día"]] is None:
            continue
        account = plain(raw[index["Nombre de la cuenta"]])
        slug = META_ACCOUNTS.get(account)
        if slug is None:
            unknown.add(account)
            continue
        campaign = plain(raw[index["Nombre de la campaña"]])
        if campaign == "":
            continue
        currency = plain(raw[index["Divisa"]]).upper()
        if currency not in ("GTQ", "USD"):
            raise RuntimeError(f"unexpected Meta currency {currency} on {account}")
        add_ad(bucket, {
            "platform": "meta",
            "cost_center_slug": slug,
            "account_name": account,
            "campaign_name": campaign,
            "day": plain(raw[index["Día"]])[:10],
            "currency": currency,
            "spend": qty(raw[index["Importe gastado (USD)"]]),
            "impressions": qty(raw[index["Impresiones"]]),
            "reach": qty(raw[index["Alcance"]]),
            "clicks": qty(raw[index["Clics (todos)"]]),
            "leads": qty(raw[index["Clientes potenciales"]]),
            "lead_forms": Decimal("0.00"),
        })
    workbook.close()
    if unknown:
        raise RuntimeError(f"unknown Meta accounts: {sorted(unknown)}")
    return bucket


def lead_forms(value: object) -> Decimal:
    total = Decimal("0.00")
    for part in plain(value).split(";"):
        folded = fold(part)
        if "clientes potenciales" not in folded or ":" not in part:
            continue
        amount = number_text(part.split(":", 1)[1])
        if amount:
            total += Decimal(amount)
    return total.quantize(Decimal("0.01"))


def load_google() -> dict[tuple[str, str, str, str], dict[str, object]]:
    path = find_file("GA -")
    raw = path.read_bytes()
    text = raw.decode("utf-16") if raw.startswith(b"\xff\xfe") else raw.decode("utf-8-sig")
    lines = text.splitlines()
    if len(lines) < 3:
        raise RuntimeError("Google Ads export is too short")
    reader = csv.DictReader(lines[2:], delimiter="\t")
    bucket: dict[tuple[str, str, str, str], dict[str, object]] = {}
    unknown: set[str] = set()
    for row in reader:
        account = plain(row.get("Nombre de la cuenta"))
        slug = GOOGLE_ACCOUNTS.get(account)
        if slug is None:
            if account:
                unknown.add(account)
            continue
        campaign = plain(row.get("Campaña"))
        day = plain(row.get("Día"))[:10]
        if campaign == "" or len(day) != 10:
            continue
        currency = plain(row.get("Código de moneda")).upper()
        if currency not in ("GTQ", "USD"):
            raise RuntimeError(f"unexpected Google currency {currency} on {account}")
        add_ad(bucket, {
            "platform": "google",
            "cost_center_slug": slug,
            "account_name": account,
            "campaign_name": campaign,
            "day": day,
            "currency": currency,
            "spend": qty(row.get("Costo")),
            "impressions": qty(row.get("Impr.")),
            "reach": None,
            "clicks": qty(row.get("Clics")),
            "leads": Decimal("0.00"),
            "lead_forms": lead_forms(row.get("Resultados")),
        })
    if unknown:
        raise RuntimeError(f"unknown Google accounts: {sorted(unknown)}")
    return bucket


def load_budget() -> list[dict[str, object]]:
    path = find_file("02")
    workbook = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict[str, object]] = []
    for sheet_name, slug in BUDGET_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(f"budget sheet {sheet_name} is missing")
        sheet = workbook[sheet_name]
        title = plain(sheet.cell(1, 1).value)
        if "2026" not in title:
            raise RuntimeError(f"{sheet_name} title has no 2026: {title}")
        month_cols: list[tuple[int, int, int]] = []
        for column in range(1, sheet.max_column + 1):
            label = plain(sheet.cell(3, column).value).upper()
            if label in MONTHS:
                month_cols.append((MONTHS[label], column, column + 1))
        if len(month_cols) != 12:
            raise RuntimeError(f"{sheet_name} has {len(month_cols)} month headers")
        for row_number in range(5, sheet.max_row + 1):
            label = plain(sheet.cell(row_number, 1).value)
            if label == "" or label.upper().startswith("TOTAL"):
                continue
            code = line_code(label)
            if code is None:
                continue
            for month, plan_column, real_column in month_cols:
                plan = money(sheet.cell(row_number, plan_column).value)
                real = money(sheet.cell(row_number, real_column).value)
                if plan is None and real is None:
                    continue
                rows.append({
                    "cost_center_slug": slug,
                    "line_code": code,
                    "month_start": f"2026-{month:02d}-01",
                    "presupuestado": None if plan is None else format(plan, "f"),
                    "real_gtq": None if real is None else format(real, "f"),
                    "source_label": label,
                })
    workbook.close()
    merged: dict[tuple[str, str, str], dict[str, object]] = {}
    for row in rows:
        key = (str(row["cost_center_slug"]), str(row["line_code"]), str(row["month_start"]))
        current = merged.get(key)
        if current is None:
            merged[key] = row
            continue
        current["presupuestado"] = add_optional(current["presupuestado"], row["presupuestado"])
        current["real_gtq"] = add_optional(current["real_gtq"], row["real_gtq"])
    return list(merged.values())


def add_optional(left: object, right: object) -> str | None:
    if left is None:
        return None if right is None else str(right)
    if right is None:
        return str(left)
    return format((Decimal(str(left)) + Decimal(str(right))).quantize(Decimal("0.01")), "f")


def meta_fields(row: dict[str, str], prefix: str) -> dict[str, str]:
    return {
        "campaign_id": plain(row.get(f"{prefix}Campaign ID")),
        "campaign_name": plain(row.get(f"{prefix}Campaign Name")),
        "adset_id": plain(row.get(f"{prefix}Adset ID")),
        "adset_name": plain(row.get(f"{prefix}Adset Name")),
        "ad_id": plain(row.get(f"{prefix}Ad ID")),
        "ad_name": plain(row.get(f"{prefix}Ad Name")),
        "form_id": plain(row.get(f"{prefix}Form ID" if prefix else "Meta Form ID")),
        "form_name": plain(row.get(f"{prefix}Form name" if prefix else "Form name")),
        "lead_id": plain(row.get(f"{prefix}Lead ID" if prefix else "Meta ID")),
    }


def person_meta(row: dict[str, str]) -> dict[str, str]:
    return {
        "campaign_id": plain(row.get("Campaign ID")),
        "campaign_name": plain(row.get("Campaign Name")),
        "adset_id": plain(row.get("Adset ID")),
        "adset_name": plain(row.get("Adset Name")),
        "ad_id": plain(row.get("Ad ID")),
        "ad_name": plain(row.get("Ad Name")),
        "form_id": plain(row.get("Meta Form ID")),
        "form_name": plain(row.get("Form name")),
        "lead_id": plain(row.get("Meta ID")),
    }


def has_meta(fields: dict[str, str]) -> bool:
    return any(fields.values())


def load_people() -> dict[str, dict[str, str]]:
    people: dict[str, dict[str, str]] = {}
    with (PIPEDRIVE / "persons.csv").open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            fields = person_meta(row)
            if fields["campaign_id"] or fields["ad_id"] or fields["lead_id"]:
                people[plain(row.get("id"))] = fields
    return people


def load_deals(people: dict[str, dict[str, str]]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    with (PIPEDRIVE / "deals.csv").open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            deal_id = plain(row.get("id"))
            if deal_id == "":
                continue
            deal_fields = meta_fields(row, "[Meta] ")
            person = people.get(plain(row.get("person_id")))
            chosen: dict[str, str] | None = None
            taken = "deal"
            if deal_fields["campaign_id"]:
                chosen = deal_fields
            elif person and person["campaign_id"]:
                chosen = person
                taken = "person"
            elif has_meta(deal_fields):
                chosen = deal_fields
            if chosen and has_meta(chosen):
                rows.append({
                    "pipedrive_deal_id": deal_id,
                    "platform": "meta",
                    "taken_from": taken,
                    **{key: value or None for key, value in chosen.items()},
                })
            tiktok_id = plain(row.get("[TikTok] Campaign ID"))
            tiktok_name = plain(row.get("[TikTok] Campaign Name"))
            if tiktok_id or tiktok_name:
                rows.append({
                    "pipedrive_deal_id": deal_id,
                    "platform": "tiktok",
                    "taken_from": "deal",
                    "campaign_id": tiktok_id or None,
                    "campaign_name": tiktok_name or None,
                    "adset_id": plain(row.get("[TikTok] Adgroup ID")) or None,
                    "adset_name": plain(row.get("[TikTok] Adgroup Name")) or None,
                    "ad_id": plain(row.get("[TikTok] Ad Id")) or None,
                    "ad_name": plain(row.get("[TikTok] Ad Name")) or None,
                    "form_id": plain(row.get("[TikTok] Form Id")) or None,
                    "form_name": plain(row.get("[TikTok] Form Name")) or None,
                    "lead_id": plain(row.get("[TikTok] Lead Id")) or None,
                })
    return rows


def decimal_text(value: object) -> str:
    return format(Decimal(str(value)).quantize(Decimal("0.01")), "f")


def main() -> None:
    check_only = "--check" in sys.argv
    print("parsing Meta", flush=True)
    ads = load_meta()
    print("parsing Google Ads", flush=True)
    ads.update(load_google())
    print("parsing budget", flush=True)
    budget = load_budget()
    print("parsing Pipedrive attribution", flush=True)
    deals = load_deals(load_people())
    excluded = sum(1 for row in ads.values() if row["campaign_name"] == EXCLUDED)
    print(
        f"ads {len(ads)} (excluded campaign rows {excluded}), budget {len(budget)}, attribution {len(deals)}",
        flush=True,
    )
    if check_only:
        return

    env = load_env()
    base = env.get("SUPABASE_URL") or env.get("NEXT_PUBLIC_SUPABASE_URL")
    key = env.get("SUPABASE_SERVICE_ROLE_KEY")
    if not base or not key:
        raise RuntimeError("SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY are required")
    api = Api(base, key)
    centers = {str(row["slug"]): str(row["id"]) for row in api.fetch_all("cost_center", "id,slug")}
    expected = set(META_ACCOUNTS.values())
    if set(centers) != expected:
        raise RuntimeError(f"cost_center slugs are {sorted(centers)}. Run migration 079 and expose schema mercadeo.")

    ad_rows: list[dict[str, object]] = []
    for row in ads.values():
        ad_rows.append({
            "id": uuid7(),
            "platform": row["platform"],
            "cost_center_id": centers[str(row["cost_center_slug"])],
            "account_name": row["account_name"],
            "campaign_name": row["campaign_name"],
            "day": row["day"],
            "currency": row["currency"],
            "spend": decimal_text(row["spend"]),
            "impressions": decimal_text(row["impressions"]),
            "reach": None if row["reach"] is None else decimal_text(row["reach"]),
            "clicks": decimal_text(row["clicks"]),
            "leads": decimal_text(row["leads"]),
            "lead_forms": decimal_text(row["lead_forms"]),
        })
    budget_rows: list[dict[str, object]] = []
    for row in budget:
        budget_rows.append({
            "id": uuid7(),
            "cost_center_id": centers[str(row["cost_center_slug"])],
            "line_code": row["line_code"],
            "month_start": row["month_start"],
            "presupuestado": row["presupuestado"],
            "real_gtq": row["real_gtq"],
        })
    attr_rows = []
    for row in deals:
        attr_rows.append({"id": uuid7(), **row})

    print("replacing facts", flush=True)
    api.delete_all("ad_day")
    api.delete_all("budget_month")
    api.delete_all("deal_attribution")
    api.insert("ad_day", ad_rows)
    api.insert("budget_month", budget_rows)
    api.insert("deal_attribution", attr_rows)
    print("done", flush=True)


if __name__ == "__main__":
    main()
